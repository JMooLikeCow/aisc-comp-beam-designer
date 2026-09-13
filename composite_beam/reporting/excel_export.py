"""Excel workbook export of a DesignResult (openpyxl)."""

from __future__ import annotations

from io import BytesIO
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from composite_beam.design_engine import DesignResult


def result_to_xlsx_bytes(result: "DesignResult") -> bytes:
    """Build a multi-sheet .xlsx workbook and return its bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    from composite_beam.reporting.summary import detailed_lines, summary_lines
    from composite_beam.units import (
        dual_force_kN,
        dual_length_mm,
        dual_moment_kNm,
        kn_to_kip,
        knm_to_kipft,
        mm_to_in,
    )

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    wrap = Alignment(wrap_text=True, vertical="top")

    def _header(ws, headers: list[str]) -> None:
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(1, col, h)
            cell.font = header_font
            cell.fill = header_fill

    # --- Summary ---
    ws = wb.active
    ws.title = "Summary"
    _header(ws, ["Item", "Value"])
    for i, line in enumerate(summary_lines(result), start=2):
        if ": " in line and not line.startswith("=") and not line.startswith("---"):
            k, _, v = line.partition(": ")
            ws.cell(i, 1, k.strip())
            ws.cell(i, 2, v.strip())
        else:
            ws.cell(i, 1, line)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 72

    # --- Checks ---
    ws2 = wb.create_sheet("Checks")
    _header(ws2, ["Check", "Demand (SI [US])", "Capacity (SI [US])", "DCR", "Pass"])
    rows = [
        (
            "Positive flexure I3",
            result.Mu_kNm,
            result.positive_moment.phiMn_kNm,
            result.DCR_flexure,
            result.pass_flexure,
        ),
    ]
    if result.negative_moment is not None:
        rows.append(
            (
                "Negative flexure Ch.F",
                abs(result.Mu_neg_kNm),
                result.negative_moment.phiMn_kNm,
                result.negative_moment.DCR,
                result.negative_moment.passes,
            )
        )
    if result.construction_LTB is not None:
        rows.append(
            (
                "Construction LTB F2",
                result.Mu_construction_kNm,
                result.construction_LTB.phiMn_kNm,
                result.construction_LTB.DCR,
                result.pass_construction,
            )
        )
    if result.interaction is not None:
        rows.append(
            (
                f"Ch.H {result.interaction.equation}",
                result.interaction.Pr_kN,
                result.interaction.Pc_kN,
                result.interaction.DCR,
                result.interaction.passes,
            )
        )
    if result.punching is not None:
        rows.append(
            (
                "Slab punching (ACI 22.6)",
                result.punching.Vu_kN,
                result.punching.phiVc_kN,
                result.punching.DCR,
                result.punching.passes,
            )
        )
    rows.append(
        (
            "Live deflection",
            result.deflection.delta_LL_mm,
            result.deflection.delta_LL_limit_mm,
            result.deflection.delta_LL_mm / result.deflection.delta_LL_limit_mm
            if result.deflection.delta_LL_limit_mm
            else 0.0,
            result.deflection.LL_OK,
        )
    )
    def _fmt_check(name: str, dem: float, cap: float) -> tuple[str, str]:
        if "deflection" in name.lower() or "Live deflection" in name:
            return dual_length_mm(float(dem), precision=1), dual_length_mm(float(cap), precision=1)
        if "punching" in name.lower() or name.startswith("Ch.H"):
            return dual_force_kN(float(dem)), dual_force_kN(float(cap))
        return dual_moment_kNm(float(dem)), dual_moment_kNm(float(cap))

    for i, (chk, dem, cap, dcr, ok) in enumerate(rows, start=2):
        dem_s, cap_s = _fmt_check(chk, dem, cap)
        ws2.cell(i, 1, chk)
        ws2.cell(i, 2, dem_s)
        ws2.cell(i, 3, cap_s)
        ws2.cell(i, 4, round(float(dcr), 3) if dcr != float("inf") else "inf")
        ws2.cell(i, 5, "PASS" if ok else "FAIL")
    for col, w in zip("ABCDE", (28, 36, 36, 10, 10)):
        ws2.column_dimensions[col].width = w

    # --- Detailed ---
    ws3 = wb.create_sheet("Detailed")
    _header(ws3, ["Calculation note"])
    for i, line in enumerate(detailed_lines(result), start=2):
        c = ws3.cell(i, 1, line)
        c.alignment = wrap
    ws3.column_dimensions["A"].width = 120

    # --- Diagram ---
    ws4 = wb.create_sheet("Diagram")
    _header(ws4, ["x_mm", "x_in", "V_kN", "V_kip", "M_kNm", "M_kipft"])
    if result.diagram is not None:
        x = result.diagram.x_mm
        V = result.diagram.V_kN
        M = result.diagram.M_kNmm / 1000.0
        for i, (xi, vi, mi) in enumerate(zip(x, V, M), start=2):
            ws4.cell(i, 1, float(xi))
            ws4.cell(i, 2, float(mm_to_in(xi)))
            ws4.cell(i, 3, float(vi))
            ws4.cell(i, 4, float(kn_to_kip(vi)))
            ws4.cell(i, 5, float(mi))
            ws4.cell(i, 6, float(knm_to_kipft(mi)))

    # --- Studs ---
    ws5 = wb.create_sheet("Studs")
    _header(ws5, ["x_mm", "zone", "row"])
    if result.stud_layout is not None:
        for i, p in enumerate(result.stud_layout.positions, start=2):
            ws5.cell(i, 1, p.x_mm)
            ws5.cell(i, 2, p.zone_name)
            ws5.cell(i, 3, p.row)

    # --- Cumulative composite action ---
    wsC = wb.create_sheet("CumulativeAction")
    _header(
        wsC,
        [
            "x_mm",
            "x_m",
            "F_req_kN",
            "F_req_kip",
            "SumQn_prov_kN",
            "SumQn_prov_kip",
            "alpha",
            "shortfall_kN",
            "status",
        ],
    )
    if result.cumulative is not None:
        for i, row in enumerate(result.cumulative.station_rows(), start=2):
            wsC.cell(i, 1, row["x_mm"])
            wsC.cell(i, 2, row["x_mm"] / 1000.0)
            wsC.cell(i, 3, row["F_req_kN"])
            wsC.cell(i, 4, float(kn_to_kip(row["F_req_kN"])))
            wsC.cell(i, 5, row["SumQn_prov_kN"])
            wsC.cell(i, 6, float(kn_to_kip(row["SumQn_prov_kN"])))
            wsC.cell(i, 7, round(row["alpha"], 4))
            wsC.cell(i, 8, row["shortfall_kN"])
            wsC.cell(i, 9, row["status"])
        # notes below table
        note_row = 3 + len(result.cumulative.station_rows())
        wsC.cell(note_row, 1, "Notes (AISC I3.2d / I8 detailing)")
        for j, n in enumerate(result.cumulative.notes):
            wsC.cell(note_row + 1 + j, 1, n)

    # --- Passing shapes ---
    ws6 = wb.create_sheet("PassingShapes")
    _header(ws6, ["Shape", "W kg/m [plf]", "phiMn (kN·m [kip·ft])", "DCR_flex", "DCR_constr", "LL_OK"])
    _kg_m_per_plf = 1.4881639437
    for i, p in enumerate(result.passing_shapes, start=2):
        ws6.cell(i, 1, p.display_name or p.designation)
        ws6.cell(i, 2, f"{p.W_lb_ft * _kg_m_per_plf:.1f} [{p.W_lb_ft:.0f}]")
        ws6.cell(i, 3, dual_moment_kNm(p.phiMn_kNm))
        ws6.cell(i, 4, round(p.DCR_flexure, 3))
        ws6.cell(i, 5, round(p.DCR_construction, 3))
        ws6.cell(i, 6, p.LL_OK)

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
